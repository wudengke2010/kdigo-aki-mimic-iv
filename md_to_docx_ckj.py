#!/usr/bin/env python3
"""Convert manuscript_ckj_submission.md to CKJ submission-formatted DOCX."""

import re
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

def add_page_number(doc):
    section = doc.sections[0]
    footer = section.footer
    footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run()
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._element.append(fldChar1)
    run2 = footer_para.add_run()
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve">PAGE</w:instrText>')
    run2._element.append(instrText)
    run3 = footer_para.add_run()
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run3._element.append(fldChar2)

def set_run_format(run, font_name='Times New Roman', font_size=12, bold=False, italic=False, superscript=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    if superscript:
        run.font.superscript = True
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)

def add_bold_italic_text(paragraph, text, font_size=12):
    bold_parts = re.split(r'(\*\*.*?\*\*)', text)
    for bp in bold_parts:
        if bp.startswith('**') and bp.endswith('**') and len(bp) > 4:
            content = bp[2:-2]
            italic_parts = re.split(r'(\*.*?\*)', content)
            for ip in italic_parts:
                if ip.startswith('*') and ip.endswith('*') and len(ip) > 2:
                    run = paragraph.add_run(ip[1:-1])
                    set_run_format(run, font_size=font_size, bold=True, italic=True)
                elif ip:
                    run = paragraph.add_run(ip)
                    set_run_format(run, font_size=font_size, bold=True)
        else:
            italic_parts = re.split(r'(\*.*?\*)', bp)
            for ip in italic_parts:
                if ip.startswith('*') and ip.endswith('*') and len(ip) > 2:
                    content = ip[1:-1]
                    run = paragraph.add_run(content)
                    set_run_format(run, font_size=font_size, italic=True)
                elif ip:
                    run = paragraph.add_run(ip)
                    set_run_format(run, font_size=font_size)

def add_formatted_text(paragraph, text, font_size=12):
    parts = re.split(r'(<sup>.*?</sup>)', text)
    for part in parts:
        if part.startswith('<sup>'):
            content = part[5:-6]
            run = paragraph.add_run(content)
            set_run_format(run, font_size=font_size, superscript=True)
        else:
            add_bold_italic_text(paragraph, part, font_size)

def parse_table(lines, start_idx):
    rows = []
    i = start_idx
    while i < len(lines) and '|' in lines[i].strip():
        line = lines[i].strip()
        if line.startswith('|'):
            line = line[1:]
        if line.endswith('|'):
            line = line[:-1]
        cells = [c.strip() for c in line.split('|')]
        rows.append(cells)
        i += 1
    if len(rows) > 1 and all(re.match(r'^[-:]+$', c.replace(' ', '')) for c in rows[1]):
        rows.pop(1)
    return rows, i

def add_table_to_doc(doc, rows, font_size=10):
    if not rows:
        return
    n_cols = len(rows[0])
    n_rows = len(rows)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            if j >= n_cols:
                break
            cell = table.rows[i].cells[j]
            cell.text = ''
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            is_header = (i == 0)
            add_formatted_text(para, cell_text, font_size=font_size)
            if is_header:
                for run in para.runs:
                    run.font.bold = True
    doc.add_paragraph()

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    style.paragraph_format.space_after = Pt(0)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    add_page_number(doc)

    i = 0
    in_references = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == '---':
            i += 1
            continue

        # Title
        if stripped.startswith('# ') and not stripped.startswith('## '):
            title_text = stripped[2:]
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(24)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = para.add_run(title_text)
            set_run_format(run, font_size=14, bold=True)
            i += 1
            continue

        # Sub-headings (###)
        if stripped.startswith('### '):
            heading_text = stripped[4:]
            para = doc.add_heading(level=2)
            para.paragraph_format.space_before = Pt(12)
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_formatted_text(para, heading_text, font_size=12)
            for run in para.runs:
                run.font.bold = True
                run.font.italic = False
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
            i += 1
            continue

        # Section headings (##)
        if stripped.startswith('## '):
            heading_text = stripped[3:]
            if 'References' in heading_text:
                in_references = True
            else:
                in_references = False

            para = doc.add_heading(level=1)
            para.paragraph_format.space_before = Pt(18)
            para.paragraph_format.space_after = Pt(12)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_formatted_text(para, heading_text, font_size=14)
            for run in para.runs:
                run.font.bold = True
                run.font.italic = False
                run.font.name = 'Times New Roman'
                run.font.size = Pt(14)
            i += 1
            continue

        # Tables
        if stripped.startswith('|'):
            rows, next_i = parse_table(lines, i)
            if rows:
                if len(rows) > 2 and len(rows[0]) > 3:
                    add_table_to_doc(doc, rows, font_size=9)
                else:
                    add_table_to_doc(doc, rows, font_size=10)
            i = next_i
            continue

        # Bullet lists
        if stripped.startswith('- '):
            bullet_text = stripped[2:]
            para = doc.add_paragraph(style='List Bullet')
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(3)
            add_formatted_text(para, bullet_text, font_size=12)
            i += 1
            continue

        # References
        if in_references and re.match(r'^\d+\.\s', stripped):
            ref_text = re.sub(r'^\d+\.\s', '', stripped)
            para = doc.add_paragraph()
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.left_indent = Inches(0.3)
            para.paragraph_format.first_line_indent = Inches(-0.3)
            add_formatted_text(para, ref_text, font_size=10)
            i += 1
            continue

        # Regular paragraph
        para = doc.add_paragraph()
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.first_line_indent = Inches(0.5)
        add_formatted_text(para, stripped, font_size=12)
        i += 1

    doc.save(docx_path)
    print(f"DOCX saved: {docx_path}")
    print(f"File size: {os.path.getsize(docx_path) / 1024:.1f} KB")

if __name__ == '__main__':
    md_path = r'C:\Users\admin\WorkBuddy\2026-07-06-13-22-19\manuscript_ckj_submission.md'
    docx_path = r'C:\Users\admin\WorkBuddy\2026-07-06-13-22-19\manuscript_ckj_submission.docx'
    convert_md_to_docx(md_path, docx_path)
