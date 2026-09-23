#!/usr/bin/env python3
"""Generate STROBE checklist DOCX for CKJ submission."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.line_spacing = 2.0

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('STROBE Checklist')
run.bold = True
run.font.size = Pt(14)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Checklist of items that should be included in reports of cohort studies').italic = True

p = doc.add_paragraph()
p.add_run('Manuscript: Association of KDIGO Creatinine-Based AKI Staging with ICU Mortality: A Dual-Cohort Retrospective Study Using MIMIC-IV v3.1 and eICU-CRD').italic = True

# STROBE items
items = [
    ('1', 'Title and abstract', 'Title indicates study design (retrospective cohort). Abstract includes design, setting, participants, main exposures, outcome, results.', 'Title page; Abstract'),
    ('2', 'Background/rationale', 'Explained the scientific background and rationale for the investigation being reported.', 'Introduction, paragraph 1'),
    ('3', 'Objectives', 'Stated specific objectives, including any prespecified hypotheses.', 'Introduction, last paragraph'),
    ('4', 'Methods: Study design', 'Presented key elements of study design early in the paper.', 'Methods, Study Design and Population'),
    ('5', 'Methods: Setting', 'Described the setting, locations, and relevant dates.', 'Methods, Study Design and Population (MIMIC-IV v3.1 2008-2022, eICU-CRD 2014-2015)'),
    ('6', 'Methods: Participants', 'Gave the eligibility criteria, and the sources and methods of selection of participants.', 'Methods, Study Design and Population; Figure 1 flowchart'),
    ('7', 'Methods: Variables', 'Clearly defined all outcomes, exposures, predictors, potential confounders, and effect modifiers.', 'Methods, Variables and Definitions; Supplementary Table S1'),
    ('8', 'Methods: Data sources', 'For each variable of interest, gave sources of data and details of methods of assessment.', 'Methods, Variables and Definitions; GitHub code with itemid mapping'),
    ('9', 'Methods: Bias', 'Described any efforts to address potential sources of bias.', 'Methods, Statistical Analysis (E-value, DAG); Discussion, Limitations'),
    ('10', 'Methods: Study size', 'Explained how the study size was arrived at.', 'Results, first paragraph; Figure 1 flowchart (N=250,540)'),
    ('11', 'Methods: Quantitative variables', 'Explained how quantitative variables were handled.', 'Methods, Variables and Definitions (creatinine ratio, APACHE IV as continuous)'),
    ('12', 'Methods: Statistical methods', 'Described all statistical methods.', 'Methods, Statistical Analysis (logistic regression, Cox, KM, NRI/IDI, TD-ROC, nomogram)'),
    ('13', 'Results: Participants', 'Reported numbers of individuals at each stage and reasons for non-participation.', 'Results, Table 1; Figure 1 flowchart'),
    ('14', 'Results: Descriptive data', 'Gave characteristics of study participants and information on exposures and potential confounders.', 'Table 1 (baseline characteristics by KDIGO stage)'),
    ('15', 'Results: Outcome data', 'Reported numbers of outcome events or summary measures.', 'Results; Table 2 (mortality by stage); Figure 2'),
    ('16', 'Results: Main results', 'Reported unadjusted estimates and confounder-adjusted estimates and their CI.', 'Table 2 (Models A-C); Figure 4 (forest plot)'),
    ('17', 'Results: Other analyses', 'Reported other analyses done (subgroups, interactions, sensitivity).', 'Results (CKD x KDIGO interaction; sepsis subgroup; MDRD sensitivity; TD-ROC; nomogram)'),
    ('18', 'Discussion: Key results', 'Summarized key results with reference to study objectives.', 'Discussion, first paragraph'),
    ('19', 'Discussion: Limitations', 'Discussed limitations of the study, considering sources of bias or imprecision.', 'Discussion, Limitations (8 items: muscle mass, urine output, ICD coding, etc.)'),
    ('20', 'Discussion: Interpretation', 'Gave a cautious overall interpretation considering objectives, limitations, and multiplicity of analyses.', 'Discussion, paragraphs 3-6'),
    ('21', 'Discussion: Generalizability', 'Discussed the generalizability (external validity) of the study results.', 'Discussion (eICU external validation; AKI paradox cross-database replication)'),
    ('22', 'Other information: Funding', 'Gave the source of funding and the role of funders.', 'Funding section'),
]

table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
hdr = table.rows[0].cells
headers = ['Item', 'Topic', 'Description', 'Location in manuscript']
for i, h in enumerate(headers):
    hdr[i].text = h
    for p in hdr[i].paragraphs:
        for r in p.runs:
            r.bold = True

for item, topic, desc, loc in items:
    row = table.add_row().cells
    row[0].text = item
    row[1].text = topic
    row[2].text = desc
    row[3].text = loc

doc.save('ckj_submission/STROBE_checklist.docx')
print('Saved: ckj_submission/STROBE_checklist.docx')
print(f'Table rows: {len(items)} STROBE items + header')
